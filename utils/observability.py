import csv
import json
import re
import time
from datetime import datetime
from html import escape
from pathlib import Path

from openpyxl import Workbook
import plotly.graph_objects as go
from plotly.offline import plot


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
EVIDENCES = ROOT / "evidencias"


def feature_folder(feature_name):
    value = (feature_name or "geral").lower()
    if "cadastro" in value:
        return "cadastro"
    if "autentica" in value or "login" in value:
        return "login"
    if "transfer" in value:
        return "transferencia"
    if "extrato" in value or "saldo" in value:
        return "extrato"
    return "geral"


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_")


def create_recorder(context, scenario):
    feature = getattr(getattr(scenario, "feature", None), "name", "")
    folder = feature_folder(feature)
    scenario_id = safe_name(scenario.name.split(" - ", 1)[0])
    evidence_dir = EVIDENCES / folder / scenario_id
    evidence_dir.mkdir(parents=True, exist_ok=True)
    recorder = ScenarioRecorder(context.page, scenario, evidence_dir, folder, scenario_id, context)
    recorder.started_at = time.perf_counter()
    context.observability = recorder
    recorder.attach()
    recorder.screenshot("01_inicio")
    return recorder


class ScenarioRecorder:
    def __init__(self, page, scenario, evidence_dir, folder, case_id, context):
        self.page = page
        self.scenario = scenario
        self.evidence_dir = evidence_dir
        self.folder = folder
        self.case_id = case_id
        self.context = context
        self.started_at = 0.0
        self.console = []
        self.network = []
        self.video_path = None
        self.final_url = ""
        self.failure_reason = ""

    def attach(self):
        self.page.on("console", self._on_console)
        self.page.on("pageerror", self._on_page_error)
        self.page.on("request", self._on_request)
        self.page.on("response", self._on_response)

    def _on_console(self, message):
        self.console.append({"type": message.type, "text": message.text, "location": message.location})

    def _on_page_error(self, error):
        self.console.append({"type": "pageerror", "text": str(error)})

    def _on_request(self, request):
        self.network.append({"method": request.method, "url": request.url, "status": None})

    def _on_response(self, response):
        for item in reversed(self.network):
            if item["url"] == response.url and item["status"] is None:
                item["status"] = response.status
                break

    def screenshot(self, name):
        try:
            self.page.screenshot(path=str(self.evidence_dir / f"{safe_name(name)}.png"), full_page=True)
        except Exception as error:
            self.console.append({"type": "screenshot_error", "text": str(error)})

    def before_step(self, step):
        self.screenshot(f"02_{safe_name(step.name)[:60]}")

    def finish(self, scenario):
        duration = time.perf_counter() - self.started_at
        failed = scenario.status == "failed"
        self.final_url = self.page.url if self.page else ""
        self.screenshot("failed" if failed else "03_sucesso")
        if failed:
            self.failure_reason = str(getattr(scenario, "exception", "Falha no cenário"))
            self._write_failure_artifacts()
        self._write_runtime_artifacts()
        return self.to_record(scenario, duration, failed)

    def _write_runtime_artifacts(self):
        console_text = "\n".join(f"[{item.get('type')}] {item.get('text')}" for item in self.console)
        (self.evidence_dir / "console.log").write_text(
            console_text,
            encoding="utf-8",
        )
        network_text = json.dumps(self.network, ensure_ascii=False, indent=2)
        (self.evidence_dir / "network.json").write_text(network_text, encoding="utf-8")
        log_dir = ROOT / "logs" / self.evidence_dir.name
        log_dir.mkdir(parents=True, exist_ok=True)
        (log_dir / "console.log").write_text(console_text, encoding="utf-8")
        (log_dir / "network.json").write_text(network_text, encoding="utf-8")
        root_logs = ROOT / "logs"
        root_logs.mkdir(parents=True, exist_ok=True)
        (root_logs / "browser_console.log").write_text(console_text, encoding="utf-8")
        (root_logs / "network.json").write_text(network_text, encoding="utf-8")
        with (root_logs / "execution.log").open("a", encoding="utf-8") as execution_log:
            execution_log.write(f"{datetime.now().isoformat(timespec='seconds')} {self.scenario.name} {self.final_url}\n")

    def _write_failure_artifacts(self):
        try:
            html = self.page.content()
        except Exception as error:
            html = f"Falha ao obter HTML: {error}"
        (self.evidence_dir / "page.html").write_text(html, encoding="utf-8")
        (self.evidence_dir / "stacktrace.log").write_text(
            f"URL: {self.final_url}\n\n{self.failure_reason}", encoding="utf-8"
        )

    def to_record(self, scenario, duration, failed):
        feature = getattr(getattr(scenario, "feature", None), "name", "")
        return {
            "test_id": scenario.name.split(" - ", 1)[0],
            "scenario": scenario.name,
            "feature": feature,
            "status": "FAIL" if failed else "PASS",
            "duration_seconds": round(duration, 3),
            "mass": "Faker / dados isolados por cenário",
            "expected_result": "Fluxo BDD concluído conforme a feature",
            "actual_result": self.failure_reason if failed else "Fluxo concluído com sucesso",
            "evidence_dir": str(self.evidence_dir.relative_to(ROOT)),
            "video": self.video_path or "",
            "url": self.final_url,
            "console_errors": sum(item.get("type") in {"error", "pageerror"} for item in self.console),
            "network_requests": len(self.network),
            "executed_at": datetime.now().isoformat(timespec="seconds"),
        }


def generate_reports(records, started_at, browser, environment):
    REPORTS.mkdir(exist_ok=True)
    for folder in ("json", "csv", "xlsx", "html", "dashboard", "graphs"):
        (REPORTS / folder).mkdir(parents=True, exist_ok=True)
    records = records or []
    total = len(records)
    passed = sum(item["status"] == "PASS" for item in records)
    failed = total - passed
    elapsed = round(time.perf_counter() - started_at, 3)
    success_rate = round((passed / total) * 100, 2) if total else 0
    average = round(sum(item["duration_seconds"] for item in records) / total, 3) if total else 0
    quality_score = max(0, round(success_rate - min(10, failed * 2)))
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "environment": environment,
        "browser": browser,
        "total_scenarios": total,
        "passed": passed,
        "failed": failed,
        "success_rate": success_rate,
        "average_duration": average,
        "total_duration": elapsed,
        "quality_score": quality_score,
        "classification": "Excellent" if quality_score >= 90 else "Good" if quality_score >= 75 else "Fair" if quality_score >= 50 else "Poor",
    }
    payload = {"summary": summary, "scenarios": records}
    (REPORTS / "json" / "execution_results.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    history_path = REPORTS / "json" / "execution_history.json"
    history = json.loads(history_path.read_text(encoding="utf-8")) if history_path.exists() else []
    history.append({"executed_at": summary["generated_at"], "passed": passed, "failed": failed, "quality_score": quality_score})
    history_path.write_text(json.dumps(history[-20:], ensure_ascii=False, indent=2), encoding="utf-8")
    _write_csv(records)
    _write_xlsx(records, summary)
    _write_plotly_graphs(records, history)
    _write_html(records, summary)
    _write_dashboard(records, summary, history)


def _write_csv(records):
    path = REPORTS / "csv" / "execution_results.csv"
    fields = list(records[0].keys()) if records else ["test_id", "scenario", "status"]
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def _xlsx_value(value):
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_xlsx(records, summary):
    book = Workbook()
    sheet = book.active
    sheet.title = "Resumo"
    for row, values in enumerate((("Métrica", "Valor"), *summary.items()), 1):
        for column, value in enumerate(values, 1):
            sheet.cell(row=row, column=column, value=_xlsx_value(value))
    details = book.create_sheet("Cenários")
    fields = list(records[0].keys()) if records else ["test_id", "scenario", "status"]
    details.append(fields)
    for record in records:
        details.append([_xlsx_value(record.get(field, "")) for field in fields])
    book.save(REPORTS / "xlsx" / "execution_results.xlsx")


def _write_plotly_graphs(records, history):
    statuses = [item["status"] for item in records]
    plot(go.Figure(go.Pie(labels=["PASS", "FAIL"], values=[statuses.count("PASS"), statuses.count("FAIL")])), filename=str(REPORTS / "graphs" / "pass_fail.html"), auto_open=False, include_plotlyjs="cdn")
    coverage = {}
    for item in records:
        key = feature_folder(item.get("feature", ""))
        coverage[key] = coverage.get(key, 0) + 1
    plot(go.Figure(go.Bar(x=list(coverage), y=list(coverage.values()))), filename=str(REPORTS / "graphs" / "feature_coverage.html"), auto_open=False, include_plotlyjs="cdn")
    plot(go.Figure(go.Bar(x=[item["test_id"] for item in records], y=[item["duration_seconds"] for item in records])), filename=str(REPORTS / "graphs" / "execution_time.html"), auto_open=False, include_plotlyjs="cdn")
    plot(go.Figure(go.Scatter(x=[item["executed_at"] for item in history], y=[item["quality_score"] for item in history], mode="lines+markers")), filename=str(REPORTS / "graphs" / "execution_history.html"), auto_open=False, include_plotlyjs="cdn")


def _artifact_links(item):
    base = "../../"
    links = []
    if item.get("evidence_dir"):
        links.append(f"<a href='{base}{escape(item['evidence_dir'].replace(chr(92), '/'))}'>Evidências</a>")
    if item.get("video"):
        links.append(f"<a href='{base}{escape(item['video'].replace(chr(92), '/'))}'>Vídeo</a>")
    if item.get("execution_log"):
        links.append(f"<a href='{base}{escape(item['execution_log'])}'>Logs</a>")
    if item.get("page_source"):
        links.append(f"<a href='{base}{escape(item['page_source'])}'>HTML</a>")
    if item.get("test_id"):
        scenario_dir = item.get("video", "").rsplit("/", 1)[0]
        if scenario_dir:
            links.append(f"<a href='{base}{escape(scenario_dir)}/test_data.json'>Massa</a>")
    return " | ".join(links)


def _scenario_gallery(item):
    parts = []
    if item.get("video"):
        video = escape(item["video"].replace(chr(92), "/"))
        parts.append(f"<video width='100%' controls preload='metadata' src='../../{video}'></video>")
    for shot in item.get("screenshots", [])[:3]:
        parts.append(f"<a href='../../{escape(shot)}' target='_blank'><img src='../../{escape(shot)}' alt='screenshot' loading='lazy'></a>")
    return "".join(parts)


def _write_html(records, summary):
    rows = "".join(
        f"<tr><td>{escape(str(item['test_id']))}</td><td>{escape(item['scenario'])}</td><td>{escape(item.get('feature', ''))}</td><td class='{item['status'].lower()}'>{item['status']}</td><td>{item['duration_seconds']}s</td><td>{escape(item.get('executed_at', ''))}</td><td>{escape(item['actual_result'])}</td><td>{_artifact_links(item)}</td></tr>"
        for item in records
    )
    table = f"<table><thead><tr><th>ID</th><th>Nome</th><th>Funcionalidade</th><th>Status</th><th>Tempo</th><th>Data</th><th>Resultado</th><th>Evidências</th></tr></thead><tbody>{rows}</tbody></table>"
    base = _html_head("Relatório BugBank")
    (REPORTS / "html" / "detailed_report.html").write_text(base + "<h1>Relatório Detalhado</h1>" + table + "</main></body></html>", encoding="utf-8")
    executive = f"<h1>Relatório Executivo</h1><section class='cards'>{_cards(summary)}</section><canvas id='executiveStatus'></canvas><script>new Chart(document.getElementById('executiveStatus'),{{type:'doughnut',data:{{labels:['PASS','FAIL'],datasets:[{{data:[{summary['passed']},{summary['failed']}],backgroundColor:['#24a148','#da1e28']}}]}}}});</script><h2>Top falhas</h2>{table}</main></body></html>"
    (REPORTS / "html" / "executive_report.html").write_text(base + executive, encoding="utf-8")


def _write_dashboard(records, summary, history):
    cards = _cards(summary)
    graphs = "".join(f"<iframe src='../graphs/{name}.html'></iframe>" for name in ("pass_fail", "feature_coverage", "execution_time", "execution_history"))
    scenarios = "".join(
        f"<details class='scenario'><summary><strong>{escape(str(item['test_id']))}</strong> — {escape(item['scenario'])} <span class='{item['status'].lower()}'>{item['status']}</span></summary><div class='scenario-body'>{_scenario_gallery(item)}<p>{_artifact_links(item)}</p><p><small>{escape(item.get('executed_at', ''))} | {item['duration_seconds']}s | console errors: {item.get('console_errors', 0)}</small></p></div></details>"
        for item in records
    )
    html = _html_head("Dashboard Executivo") + f"<h1>Dashboard Executivo</h1><p>Execução: {summary['generated_at']} | Ambiente: {summary['environment']} | Browser: {summary['browser']}</p>{cards}<h2>Quality Score: {summary['quality_score']}/100 ({summary['classification']})</h2><section class='graphs'>{graphs}</section><h2>Cenários</h2>{scenarios}</main></body></html>"
    (REPORTS / "dashboard" / "index.html").write_text(html, encoding="utf-8")


def _cards(summary):
    return "".join(f"<div class='card'><small>{escape(str(key))}</small><strong>{escape(str(value))}</strong></div>" for key, value in summary.items())


def _html_head(title):
    return f"<!doctype html><html lang='pt-br'><head><meta charset='utf-8'><title>{escape(title)}</title><script src='https://cdn.jsdelivr.net/npm/chart.js'></script><style>body{{font-family:Arial,sans-serif;background:#f4f6f8;color:#16212b;margin:32px}}h1{{color:#0f62fe}}.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:20px 0}}.card{{background:white;border-radius:8px;padding:16px;box-shadow:0 2px 8px #0001}}.card strong{{display:block;font-size:24px;margin-top:8px}}table{{width:100%;border-collapse:collapse;background:white}}th,td{{padding:10px;border-bottom:1px solid #ddd;text-align:left}}.pass{{color:#198038;font-weight:bold}}.fail{{color:#da1e28;font-weight:bold}}.charts{{display:grid;grid-template-columns:1fr 1fr;gap:24px;background:white;padding:20px}}.graphs{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}.graphs iframe{{width:100%;height:380px;border:0;background:white;border-radius:8px}}.scenario{{background:white;border-radius:8px;margin:10px 0;padding:12px 16px;box-shadow:0 2px 8px #0001}}.scenario summary{{cursor:pointer}}.scenario-body{{margin-top:12px}}.scenario-body img{{max-width:32%%;border-radius:6px;margin:4px 4px 4px 0;border:1px solid #ddd}}</style></head><body><main>"
